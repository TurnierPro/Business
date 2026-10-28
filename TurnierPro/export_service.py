"""Exporte aus der App: Excel (openpyxl) und PDF (reportlab), immer hell und im TurnierPro-Grün.

Bereiche: vertrieb, aufgaben, projekte. Dazu eine komplette Sicherung als JSON.
"""

from __future__ import annotations

import io
import json
from datetime import date, datetime

from speicher import Speicher

GRUEN = "1E7B4B"
MINT = "E5EEE8"
OFFEN_NICHT = {"Gewonnen", "Verloren / kein Interesse"}
STATUS = ["Erstkontakt erfolgt", "In Kontakt / Nachfassen", "Interesse bekundet", "Demo / Termin vereinbart",
          "Angebot gesendet", "Gewonnen", "Verloren / kein Interesse", "Pausiert (später aufgreifen)"]
WDH = {"": "Einmalig", "daily": "Täglich", "weekdays": "Werktags (Mo–Fr)", "weekly": "Wöchentlich", "biweekly": "Alle 2 Wochen",
       "monthly": "Monatlich", "quarterly": "Vierteljährlich", "yearly": "Jährlich"}
PROJEKT_STATUS = ["Idee", "Geplant", "In Arbeit", "Erledigt"]


def _de(iso: str) -> str:
    return f"{iso[8:10]}.{iso[5:7]}.{iso[0:4]}" if iso and len(iso) >= 10 else ""


def _eur(n) -> str:
    return f"{float(n or 0):,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


# ── Modell: Blätter mit Tabellen ────────────────────────────────────────
def _modell(sp: Speicher, bereich: str) -> dict:
    sp.aktualisieren()
    d = sp.alles()
    team = {t["id"]: t.get("name", "") for t in d["team"]}
    name = lambda i: team.get(i, "") if i else ""  # noqa: E731
    heute = date.today().isoformat()

    if bereich == "vertrieb":
        vereine = sorted(d["vereine"], key=lambda v: str(v.get("name", "")).lower())
        vname = {v["id"]: v.get("name", "") for v in vereine}
        offen = [v for v in vereine if v.get("status") not in OFFEN_NICHT]
        gew = [v for v in vereine if v.get("status") == "Gewonnen"]
        faellig = [v for v in offen if v.get("wiedervorlage") and v["wiedervorlage"] <= heute]
        kpi = [("Vereine gesamt", len(vereine)), ("Offene Vereine", len(offen)), ("Gewonnen", len(gew)),
               ("Offenes Potenzial", _eur(sum(float(v.get("potenzial") or 0) for v in offen))),
               ("Gewonnenes Volumen", _eur(sum(float(v.get("potenzial") or 0) for v in gew))),
               ("Wiedervorlagen fällig", len(faellig)), ("Aktivitäten", len(d["aktivitaeten"]))]
        pipeline = [[s, sum(1 for v in vereine if v.get("status") == s),
                     _eur(sum(float(v.get("potenzial") or 0) for v in vereine if v.get("status") == s))] for s in STATUS]
        termin = lambda v: "offen" if v.get("terminOffen") and not v.get("terminVon") else _de(v.get("terminVon", "")) + (f" – {_de(v['terminBis'])}" if v.get("terminBis") and v["terminBis"] != v.get("terminVon") else "")  # noqa: E731
        wv = sorted([v for v in offen if v.get("wiedervorlage")], key=lambda v: v["wiedervorlage"])
        akts = sorted(d["aktivitaeten"], key=lambda a: (a.get("datum") or "", a.get("angelegt") or 0), reverse=True)
        return {"titel": "Vertrieb", "blaetter": [
            {"name": "Übersicht", "kpi": kpi, "tabellen": [("Pipeline nach Status", ["Status", "Vereine", "Potenzial"], pipeline)]},
            {"name": "Vereine", "tabellen": [("", ["Verein", "Ort", "Bundesland", "Landesverband", "Ansprechperson", "Funktion", "Telefon", "E-Mail",
                                                    "Betreuer", "Veranstaltung", "Art", "Termin", "Status", "Potenzial", "Prio", "Wiedervorlage", "Nächster Schritt"],
                                              [[v.get("name", ""), " ".join(x for x in (v.get("plz", ""), v.get("ort", "")) if x), v.get("bundesland", ""), v.get("verband", ""),
                                                v.get("person", ""), v.get("funktion", ""), v.get("telefon", ""), v.get("email", ""), name(v.get("betreuer")),
                                                v.get("veranstaltung", ""), v.get("art", ""), termin(v), v.get("status", ""), float(v.get("potenzial") or 0),
                                                v.get("prio", ""), _de(v.get("wiedervorlage", "")), v.get("naechsterSchritt", "")] for v in vereine])],
             "pdf_spalten": [0, 1, 4, 8, 12, 13, 15]},
            {"name": "Wiedervorlagen", "tabellen": [("", ["Fällig am", "Verein", "Ansprechperson", "Telefon", "Betreuer", "Status", "Nächster Schritt"],
                                                     [[_de(v["wiedervorlage"]), v.get("name", ""), v.get("person", ""), v.get("telefon", ""), name(v.get("betreuer")),
                                                       v.get("status", ""), v.get("naechsterSchritt", "")] for v in wv])]},
            {"name": "Aktivitäten", "tabellen": [("", ["Datum", "Verein", "Kontaktart", "Ergebnis", "Mitarbeiter", "Notiz"],
                                                  [[_de(a.get("datum", "")), vname.get(a.get("vereinId"), "(gelöscht)"), a.get("kanal", ""), a.get("ergebnis", ""),
                                                    name(a.get("von")), a.get("notiz", "")] for a in akts])]},
        ]}

    if bereich == "aufgaben":
        offen = sorted([t for t in d["aufgaben"] if not t.get("erledigt")], key=lambda t: t.get("faellig") or "9999")
        erl = sorted([t for t in d["aufgaben"] if t.get("erledigt")], key=lambda t: t.get("erledigtAm") or "", reverse=True)
        zeile = lambda t: [t.get("titel", ""), name(t.get("zustaendig")) or "Nicht zugewiesen", _de(t.get("faellig", "")), WDH.get(t.get("wdh", ""), ""),  # noqa: E731
                           "Hoch" if t.get("prio") == "high" else "Normal",
                           "Erledigt" if t.get("erledigt") else ("Überfällig" if t.get("faellig") and t["faellig"] < heute else "Offen"), t.get("notizen", "")]
        kopf = ["Aufgabe", "Zugewiesen an", "Erledigt bis", "Wiederholung", "Priorität", "Status", "Notizen"]
        return {"titel": "Aufgaben", "blaetter": [{"name": "Aufgaben", "tabellen": [
            (f"Offene Aufgaben ({len(offen)})", kopf, [zeile(t) for t in offen]),
            (f"Erledigte Aufgaben ({len(erl)})", kopf, [zeile(t) for t in erl])]}]}

    if bereich == "projekte":
        projekte = d["projekte"]
        ptitel = {p["id"]: p.get("titel", "") for p in projekte}
        komm = sorted(d["kommentare"], key=lambda c: (ptitel.get(c.get("projektId"), ""), c.get("zeit") or 0))
        kanz = {}
        for c in d["kommentare"]:
            kanz[c.get("projektId")] = kanz.get(c.get("projektId"), 0) + 1
        tabellen = []
        for st in PROJEKT_STATUS:
            l = [p for p in projekte if (p.get("status") if p.get("status") in PROJEKT_STATUS else "Idee") == st]
            if l:
                tabellen.append((f"{st} ({len(l)})", ["Projekt", "Priorität", "Verantwortlich", "Kommentare", "Beschreibung / Notizen"],
                                 [[p.get("titel", ""), {"hoch": "Hoch", "niedrig": "Niedrig"}.get(p.get("prio"), "Normal"),
                                   ", ".join(name(i) for i in (p.get("verantwortlich") or [])) or "Nicht zugewiesen", kanz.get(p["id"], 0), p.get("notizen", "")] for p in l]))
        return {"titel": "Projekte", "blaetter": [
            {"name": "Projekte", "tabellen": tabellen or [("", ["Projekt"], [["Noch keine Projekte"]])]},
            {"name": "Kommentare", "tabellen": [("", ["Projekt", "Zeitpunkt", "Von", "Kommentar"],
                                                 [[ptitel.get(c.get("projektId"), ""), datetime.fromtimestamp((c.get("zeit") or 0) / 1000).strftime("%d.%m.%Y %H:%M") if c.get("zeit") else "",
                                                   name(c.get("von")), c.get("text", "")] for c in komm])]},
        ]}
    raise ValueError("Unbekannter Bereich.")


# ── Excel ───────────────────────────────────────────────────────────────
def excel(sp: Speicher, bereich: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    m = _modell(sp, bereich)
    wb = Workbook()
    wb.remove(wb.active)
    for b in m["blaetter"]:
        ws = wb.create_sheet(b["name"][:31])
        ws.sheet_view.showGridLines = False
        ws["A1"] = f"TurnierPro · {m['titel']} · {b['name']}"
        ws["A1"].font = Font(bold=True, size=14, color=GRUEN)
        ws["A2"] = f"Stand {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        ws["A2"].font = Font(size=9, color="6B7A71")
        r = 4
        breiten: dict[int, int] = {}
        for k, v in b.get("kpi", []):
            ws.cell(r, 1, k).font = Font(bold=True)
            ws.cell(r, 2, v)
            r += 1
        if b.get("kpi"):
            r += 1
        for titel, kopf, zeilen in b["tabellen"]:
            if titel:
                ws.cell(r, 1, titel).font = Font(bold=True, size=11, color=GRUEN)
                r += 1
            kopf_zeile = r
            for j, h in enumerate(kopf, 1):
                c = ws.cell(r, j, h)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=GRUEN)
                c.alignment = Alignment(vertical="center")
                breiten[j] = max(breiten.get(j, 0), len(h) + 2)
            r += 1
            for i, zeile in enumerate(zeilen):
                for j, w in enumerate(zeile, 1):
                    c = ws.cell(r, j, w)
                    c.alignment = Alignment(vertical="top", wrap_text=isinstance(w, str) and len(w) > 50)
                    if isinstance(w, float):
                        c.number_format = '#,##0.00 "€"'
                    if i % 2:
                        c.fill = PatternFill("solid", fgColor="F7FAF8")
                    breiten[j] = min(60, max(breiten.get(j, 0), len(str(w)) + 2))
                r += 1
            if zeilen and len(b["tabellen"]) == 1:
                ws.auto_filter.ref = f"A{kopf_zeile}:{get_column_letter(len(kopf))}{r - 1}"
                ws.freeze_panes = f"A{kopf_zeile + 1}"
            r += 1
        for j, w in breiten.items():
            ws.column_dimensions[get_column_letter(j)].width = max(10, w)
    puffer = io.BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


# ── PDF ─────────────────────────────────────────────────────────────────
def pdf(sp: Speicher, bereich: str, logo_pfad=None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    m = _modell(sp, bereich)
    gruen, mint, grau = colors.HexColor("#" + GRUEN), colors.HexColor("#" + MINT), colors.HexColor("#6B7A71")
    st_titel = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=17, textColor=colors.HexColor("#1A1A1A"), leading=21)
    st_sub = ParagraphStyle("s", fontName="Helvetica", fontSize=9, textColor=grau)
    st_h2 = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=11.5, textColor=gruen, spaceBefore=10, spaceAfter=5)
    st_z = ParagraphStyle("z", fontName="Helvetica", fontSize=7.8, leading=9.6)
    st_k = ParagraphStyle("k", fontName="Helvetica-Bold", fontSize=7.8, leading=9.6, textColor=colors.white)
    seite = landscape(A4)
    puffer = io.BytesIO()
    doc = SimpleDocTemplate(puffer, pagesize=seite, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
                            title=f"TurnierPro · {m['titel']}", author="TurnierPro")
    breite = seite[0] - 28 * mm

    def fuss(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(grau)
        canvas.drawString(14 * mm, 8 * mm, f"TurnierPro UG (haftungsbeschränkt) · {m['titel']} · Stand {datetime.now().strftime('%d.%m.%Y')}")
        canvas.drawRightString(seite[0] - 14 * mm, 8 * mm, f"Seite {_doc.page}")
        canvas.restoreState()

    teile = []
    if logo_pfad and logo_pfad.is_file():
        teile += [Image(str(logo_pfad), width=44 * mm, height=44 * mm * 165 / 626, hAlign="LEFT"), Spacer(1, 4 * mm)]
    teile += [Paragraph(m["titel"], st_titel), Paragraph(f"Stand {datetime.now().strftime('%d.%m.%Y, %H:%M')} Uhr", st_sub), Spacer(1, 5 * mm)]
    for i, b in enumerate(m["blaetter"]):
        if i and b["name"] in ("Vereine", "Aktivitäten", "Kommentare"):
            teile.append(PageBreak())
        if b.get("kpi"):
            zellen = [[Paragraph(f"<font size=7 color='#6B7A71'>{k.upper()}</font><br/><b>{v}</b>", st_z) for k, v in b["kpi"]]]
            t = Table(zellen, colWidths=[breite / len(b["kpi"])] * len(b["kpi"]))
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), mint), ("BOX", (0, 0), (-1, -1), 0.5, colors.white),
                                   ("INNERGRID", (0, 0), (-1, -1), 2, colors.white), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            teile += [t, Spacer(1, 4 * mm)]
        for titel, kopf, zeilen in b["tabellen"]:
            spalten = b.get("pdf_spalten") or list(range(len(kopf)))
            teile.append(Paragraph(titel or b["name"], st_h2))
            if not zeilen:
                teile.append(Paragraph("Keine Einträge.", st_sub))
                continue
            daten = [[Paragraph(str(kopf[j]), st_k) for j in spalten]]
            for z in zeilen:
                daten.append([Paragraph(_eur(z[j]) if isinstance(z[j], float) else str(z[j]).replace("&", "&amp;").replace("<", "&lt;"), st_z) for j in spalten])
            t = Table(daten, repeatRows=1, colWidths=[breite / len(spalten)] * len(spalten))
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), gruen), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAF8")]),
                                   ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#E1E8E3")),
                                   ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            teile.append(t)
    doc.build(teile, onFirstPage=fuss, onLaterPages=fuss)
    return puffer.getvalue()


def sicherung(sp: Speicher) -> bytes:
    sp.aktualisieren()
    inhalt = {"app": "TurnierPro", "art": "sicherung", "erstellt": datetime.now().isoformat(timespec="seconds"), "daten": sp.alles()}
    return json.dumps(inhalt, ensure_ascii=False, indent=1).encode("utf-8")
