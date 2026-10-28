"""Import der Artefakt-Exporte, Export als Excel/PDF/JSON und Schutz der Schnittstelle."""
import io
from datetime import datetime

import openpyxl
import pytest

import export_service
import import_service
from speicher import Speicher


def artefakt_vertrieb() -> bytes:
    """Nachbau des Excel-Exports aus dem Artefakt (Titelzeilen, dann Kopfzeile)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Übersicht"
    ws.append(["Vertrieb · Übersicht"])
    ws.append(["Stand 01.10.2026"])
    ws.append(["Vereine gesamt", 2])
    v = wb.create_sheet("Vereine")
    v.append(["Vereine und Pipeline"]); v.append(["2 Vereine"]); v.append(["Stand: … · TurnierPro"]); v.append([])
    v.append(["Verein", "Landesverband", "Bundesland", "Straße", "PLZ", "Ort", "Ansprechperson", "Funktion", "Telefon", "E-Mail",
              "Entscheider", "Betreuer", "Veranstaltung", "Art", "Termin von", "Termin bis", "Datum unbekannt", "Teilnehmer",
              "Aktuelle Lösung", "Status", "Potenzial", "Prio", "Letzter Kontakt", "Wiedervorlage", "Nächster Schritt"])
    v.append(["SV Muster", "Badischer Fußballverband", "Baden-Württemberg", "Weg 1", "76133", "Karlsruhe", "Anna Test", "Vorstand", "0721 1",
              "a@sv.example", "Ja", "Max Steiner", "Muster-Cup", "Turnier", datetime(2026, 6, 1), datetime(2026, 6, 2), "", 48,
              "Excel / Papier", "Interesse bekundet", 850.0, "A", datetime(2026, 9, 1), datetime(2026, 10, 30), "Demo anbieten"])
    v.append(["FC Probe", "", "Bayern", "", "", "Augsburg", "", "", "", "", "", "Felix Steiner", "", "", None, None, "Ja", None,
              "", "Gewonnen", 0, "", None, None, ""])
    w = wb.create_sheet("Wiedervorlagen")
    w.append(["Wiedervorlagen"]); w.append([]); w.append(["Fällig am", "Zustand", "Verein", "Ansprechperson", "Telefon", "Betreuer", "Status", "Nächster Schritt"])
    a = wb.create_sheet("Aktivitäten")
    a.append(["Aktivitäten"]); a.append([])
    a.append(["Datum", "Verein", "Kontaktart", "Ergebnis", "Mitarbeiter", "Notiz"])
    a.append([datetime(2026, 9, 1), "SV Muster", "Telefon (Cold Call)", "Gespräch geführt, Interesse", "Max Steiner", "Will Demo"])
    a.append([datetime(2026, 9, 3), "FC Probe", "E-Mail", "Abschluss (per E-Mail bestätigt)", "Felix Steiner", ""])
    a.append([datetime(2026, 9, 4), "Gibt es nicht", "E-Mail", "Keine Rückmeldung", "", ""])
    puffer = io.BytesIO(); wb.save(puffer)
    return puffer.getvalue()


def artefakt_aufgaben_projekte() -> tuple[bytes, bytes]:
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Aufgaben"
    ws.append(["Aufgaben und Umsatzziel"]); ws.append([])
    ws.append(["Aufgabe", "Zugewiesen an", "Erledigt bis", "Wiederholung", "Priorität", "Status", "Notizen"])
    ws.append(["Offene Aufgaben (1)"])
    ws.append(["Belege sammeln", "Max Steiner", datetime(2026, 10, 31), "Monatlich", "Hoch", "Offen", ""])
    ws.append(["Erledigte Aufgaben (1)"])
    ws.append(["Flyer drucken", "Nicht zugewiesen", None, "Einmalig", "Normal", "Erledigt", "fertig"])
    p1 = io.BytesIO(); wb.save(p1)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Projekte"
    ws.append(["Projekte und Ideen"]); ws.append([])
    ws.append(["Projekt", "Status", "Priorität", "Zugewiesen an", "Kommentare", "Beschreibung / Notizen"])
    ws.append(["Idee (1)"])
    ws.append(["Messe-Aufsteller", "Idee", "Hoch", "Max Steiner, Felix Steiner", 1, "Für die Hallensaison"])
    k = wb.create_sheet("Kommentare")
    k.append(["Kommentare"]); k.append([])
    k.append(["Projekt", "Zeitpunkt", "Von", "Kommentar / aktueller Stand"])
    k.append(["Messe-Aufsteller", datetime(2026, 9, 2, 10, 30), "Felix Steiner", "Angebot von der Druckerei liegt vor"])
    p2 = io.BytesIO(); wb.save(p2)
    return p1.getvalue(), p2.getvalue()


@pytest.fixture
def sp(tmp_path):
    return Speicher(tmp_path / "daten", "test")


def test_vertrieb_import_pruefen_und_ausfuehren(sp):
    datei = artefakt_vertrieb()
    vorschau = import_service.importiere(sp, "artefakt", "Vertrieb.xlsx", datei, ausfuehren=False)
    assert vorschau["sammlungen"]["vereine"]["neu"] == 2
    assert vorschau["sammlungen"]["aktivitaeten"]["neu"] == 2
    assert vorschau["sammlungen"]["team"]["neu"] == 2
    assert any("übersprungen" in h for h in vorschau["hinweise"])
    assert sp.liste("vereine") == []  # Prüfen schreibt nichts

    import_service.importiere(sp, "artefakt", "Vertrieb.xlsx", datei, ausfuehren=True)
    vereine = {v["name"]: v for v in sp.liste("vereine")}
    sv = vereine["SV Muster"]
    assert sv["terminVon"] == "2026-06-01" and sv["potenzial"] == 850 and sv["groesse"] == 48
    assert vereine["FC Probe"]["terminOffen"] is True
    team = {t["id"]: t["name"] for t in sp.liste("team")}
    assert team[sv["betreuer"]] == "Max Steiner"
    akt = {a["notiz"]: a for a in sp.liste("aktivitaeten")}
    # Anzeigetext je Kontaktart wird zum gespeicherten Grundwert
    assert akt["Will Demo"]["ergebnis"] == "Rückmeldung und Interesse"
    assert akt[""]["ergebnis"] == "Abschluss"

    # Zweiter Import derselben Datei: alles doppelt, nichts Neues
    nochmal = import_service.importiere(sp, "artefakt", "Vertrieb.xlsx", datei, ausfuehren=False)
    assert nochmal["neu_gesamt"] == 0


def test_aufgaben_und_projekte_import(sp):
    aufg, proj = artefakt_aufgaben_projekte()
    import_service.importiere(sp, "artefakt", "Aufgaben.xlsx", aufg, ausfuehren=True)
    import_service.importiere(sp, "artefakt", "Projekte.xlsx", proj, ausfuehren=True)
    aufgaben = {a["titel"]: a for a in sp.liste("aufgaben")}
    assert aufgaben["Belege sammeln"]["wdh"] == "monthly" and aufgaben["Belege sammeln"]["prio"] == "high"
    assert aufgaben["Flyer drucken"]["erledigt"] is True and aufgaben["Flyer drucken"]["zustaendig"] == ""
    assert len(sp.liste("aufgaben")) == 2
    [p] = sp.liste("projekte")
    assert p["prio"] == "hoch" and len(p["verantwortlich"]) == 2
    [k] = sp.liste("kommentare")
    assert k["projektId"] == p["id"] and k["zeit"] > 0


def test_termine_und_besuche_csv(sp):
    csv_t = "Titel;Datum;Ende;Ort;Art;Veranstalter;URL\nHerbst-Cup;01.11.2026;02.11.2026;Karlsruhe;Turnier;SV Muster;https://www.turnierpro.de/x\n".encode()
    erg = import_service.importiere(sp, "termine", "termine.csv", csv_t, ausfuehren=True)
    assert erg["geschrieben"] == 1
    [t] = sp.liste("termine")
    assert t["datum"] == "2026-11-01" and t["art"] == "turnier"
    csv_b = "Datum,Besucher,Klicks\n2026-10-01,120,300\n2026-10-02,90,210\n".encode()
    import_service.importiere(sp, "besuche", "b.csv", csv_b, ausfuehren=True)
    assert len(sp.liste("besuche")) == 2
    erneut = import_service.importiere(sp, "besuche", "b.csv", csv_b, ausfuehren=False)
    assert erneut["sammlungen"]["besuche"]["doppelt"] == 2


def test_export_und_sicherung_rundreise(sp, tmp_path):
    import_service.importiere(sp, "artefakt", "Vertrieb.xlsx", artefakt_vertrieb(), ausfuehren=True)
    xlsx = export_service.excel(sp, "vertrieb")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Übersicht", "Vereine", "Wiedervorlagen", "Aktivitäten"]
    assert export_service.pdf(sp, "vertrieb").startswith(b"%PDF")
    assert export_service.pdf(sp, "aufgaben").startswith(b"%PDF")
    assert export_service.excel(sp, "projekte")[:2] == b"PK"
    # Sicherung in einen leeren Ordner zurückspielen
    neu = Speicher(tmp_path / "anders", "b")
    erg = import_service.importiere(neu, "sicherung", "s.json", export_service.sicherung(sp), ausfuehren=True)
    assert erg["geschrieben"] == len(sp.liste("vereine")) + len(sp.liste("aktivitaeten")) + len(sp.liste("team"))
    assert {v["id"] for v in neu.liste("vereine")} == {v["id"] for v in sp.liste("vereine")}
